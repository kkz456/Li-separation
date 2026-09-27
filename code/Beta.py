#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
从 FREQ_MASS=6 和 FREQ_MASS=7 的 VASP 频率/本征矢输出中计算
Gamma 点谐振近似的 Li 同位素 beta^(7/6)，并按材料 ID 合并到
前面 CAVD bottleneck Excel 中。

本版本对虚频的处理原则：
1. 所有标记为 fi 或 f/i 的虚频，不论大小，全部不参与 beta 计算；
2. 先用本征矢对 MASS=6 与 MASS=7 的全部模式进行配对；
3. 只保留“两套计算中都为正实频”的配对模式；
4. 低于声学零模阈值的正实频配对也排除；
5. 虚频仅在 Excel 中记录数量，不再令材料失去相关性分析资格。

推荐目录结构：
root/
├── calculate_beta_merge_bottleneck.py
├── FREQ_MASS=6/
├── FREQ_MASS=7/
├── crystal_structure/
└── cavd_Radical_Voronoi_bottleneck已补全.xlsx

运行：
python calculate_beta_merge_bottleneck.py \
  --root . \
  --bottleneck "cavd_Radical_Voronoi_bottleneck已补全.xlsx" \
  --temperature 300

输出：
  cavd_bottleneck_beta_忽略全部虚频_300K汇总.xlsx

依赖：
  python -m pip install numpy openpyxl
可选：
  python -m pip install scipy
如果安装 scipy，模式匹配使用 Hungarian 算法；否则自动使用贪心匹配。
"""

from __future__ import annotations

import argparse
import math
import re
import shutil
import sys
from copy import copy
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

import numpy as np
from openpyxl import load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

try:
    from scipy.optimize import linear_sum_assignment
except Exception:
    linear_sum_assignment = None


# -----------------------------
# 物理常数和默认参数
# -----------------------------
H_EV_S = 4.135667696e-15
KB_EV_K = 8.617333262145e-5
THZ_TO_CM1 = 33.35640951981521

DEFAULT_TEMPERATURE = 300.0
DEFAULT_ACOUSTIC_CUTOFF_CM1 = 5.0


@dataclass
class Mode:
    index: int
    imaginary: bool
    freq_thz: float
    freq_cm1: float
    eigenvector: np.ndarray | None


@dataclass
class StructureInfo:
    path: Path
    symbols: list[str]
    counts: list[int]
    atom_symbols: list[str]
    n_atoms: int
    n_li: int


# -----------------------------
# 材料 ID 对齐
# -----------------------------
def normalized_material_id(name: str) -> str:
    """把 mp-1020059、1020059、mp-6521_LLTO、6521_LLTO 对齐。"""
    text = Path(str(name)).stem.strip().lower()
    text = text.replace(" ", "")
    if text.startswith("mp-"):
        text = text[3:]
    return text


def build_file_index(directory: Path, suffixes: Iterable[str]) -> dict[str, Path]:
    suffix_set = {suffix.lower() for suffix in suffixes}
    result: dict[str, Path] = {}
    if not directory.exists():
        return result

    for path in sorted(directory.iterdir()):
        if path.is_file() and path.suffix.lower() in suffix_set:
            key = normalized_material_id(path.stem)
            if key in result:
                raise ValueError(
                    f"目录 {directory} 中有重复规范 ID：{key}\n"
                    f"  {result[key]}\n  {path}"
                )
            result[key] = path
    return result


# -----------------------------
# 读取 VASP POSCAR/vasp
# -----------------------------
def is_integer_token(token: str) -> bool:
    try:
        int(token)
        return True
    except ValueError:
        return False


def read_vasp_structure_info(path: Path) -> StructureInfo:
    lines = [line.strip() for line in path.read_text(encoding="utf-8", errors="ignore").splitlines()]
    if len(lines) < 7:
        raise ValueError(f"VASP 文件行数不足：{path}")

    line6 = lines[5].split()
    line7 = lines[6].split()

    if line6 and all(is_integer_token(token) for token in line6):
        raise ValueError(
            f"{path} 看起来是 VASP4 格式，没有元素符号行，无法自动统计 Li。"
        )

    symbols = line6
    if not line7 or not all(is_integer_token(token) for token in line7):
        raise ValueError(f"无法读取元素数目行：{path}")

    counts = [int(token) for token in line7]
    if len(symbols) != len(counts):
        raise ValueError(f"元素符号和计数长度不一致：{path}")

    atom_symbols: list[str] = []
    for symbol, count in zip(symbols, counts):
        atom_symbols.extend([symbol] * count)

    n_li = sum(count for symbol, count in zip(symbols, counts) if symbol == "Li")
    if n_li <= 0:
        raise ValueError(f"结构中没有 Li：{path}")

    return StructureInfo(
        path=path,
        symbols=symbols,
        counts=counts,
        atom_symbols=atom_symbols,
        n_atoms=sum(counts),
        n_li=n_li,
    )


# -----------------------------
# 解析 VASP 频率与本征矢
# -----------------------------
MODE_HEADER_RE = re.compile(
    r"^\s*(?P<index>\d+)\s+(?P<kind>f/i|fi|f)\s*=\s*"
    r"(?P<freq>[+-]?(?:\d+(?:\.\d*)?|\.\d+))\s*THz",
    flags=re.IGNORECASE,
)
FLOAT_RE = re.compile(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[Ee][+-]?\d+)?")


def parse_frequency_file(path: Path) -> list[Mode]:
    lines = path.read_text(encoding="utf-8", errors="ignore").splitlines()
    modes: list[Mode] = []
    i = 0

    while i < len(lines):
        match = MODE_HEADER_RE.match(lines[i])
        if not match:
            i += 1
            continue

        mode_index = int(match.group("index"))
        kind = match.group("kind").lower()
        imaginary = kind in {"fi", "f/i"}
        freq_thz = abs(float(match.group("freq")))

        vectors: list[list[float]] = []
        j = i + 1
        while j < len(lines):
            if MODE_HEADER_RE.match(lines[j]):
                break
            numbers = FLOAT_RE.findall(lines[j])
            # 原子行固定为 X Y Z dx dy dz，共六个数。
            if len(numbers) == 6:
                values = [float(number) for number in numbers]
                vectors.append(values[3:6])
            j += 1

        eigenvector: np.ndarray | None
        if vectors:
            eigenvector = np.asarray(vectors, dtype=float)
        else:
            eigenvector = None

        modes.append(
            Mode(
                index=mode_index,
                imaginary=imaginary,
                freq_thz=freq_thz,
                freq_cm1=freq_thz * THZ_TO_CM1,
                eigenvector=eigenvector,
            )
        )
        i = j

    if not modes:
        raise ValueError(f"没有从文件中解析出振动模式：{path}")

    indices = [mode.index for mode in modes]
    if len(indices) != len(set(indices)):
        raise ValueError(f"模式编号重复：{path}")

    return modes


# -----------------------------
# 模式筛选与配对
# -----------------------------
def imaginary_modes(modes: list[Mode]) -> list[Mode]:
    """返回文件中全部虚频，不设置大小阈值。"""
    return [mode for mode in modes if mode.imaginary]


def select_usable_pairs(
    all_pairs: list[tuple[Mode, Mode, float]],
    acoustic_cutoff_cm1: float,
) -> tuple[
    list[tuple[Mode, Mode, float]],
    int,
    int,
]:
    """
    从已经按本征矢匹配的全部模式对中选择 beta 模式。

    规则：
    - 任意一侧为虚频：整对排除；
    - 两侧都为实频，但任意一侧不高于声学零模阈值：整对排除；
    - 只保留两侧均为正实频且均高于阈值的模式对。

    返回：可用模式对、因虚频排除的配对数、因低频排除的配对数。
    """
    usable: list[tuple[Mode, Mode, float]] = []
    excluded_imaginary = 0
    excluded_low_frequency = 0

    for mode6, mode7, overlap in all_pairs:
        if mode6.imaginary or mode7.imaginary:
            excluded_imaginary += 1
            continue

        if (
            mode6.freq_cm1 <= acoustic_cutoff_cm1
            or mode7.freq_cm1 <= acoustic_cutoff_cm1
        ):
            excluded_low_frequency += 1
            continue

        usable.append((mode6, mode7, overlap))

    return usable, excluded_imaginary, excluded_low_frequency


def vector_overlap(mode6: Mode, mode7: Mode) -> float:
    if mode6.eigenvector is None or mode7.eigenvector is None:
        return 0.0
    if mode6.eigenvector.shape != mode7.eigenvector.shape:
        return 0.0

    v6 = mode6.eigenvector.reshape(-1)
    v7 = mode7.eigenvector.reshape(-1)
    norm = np.linalg.norm(v6) * np.linalg.norm(v7)
    if norm <= 0:
        return 0.0
    return float(abs(np.dot(v6, v7)) / norm)


def pair_modes(modes6: list[Mode], modes7: list[Mode]) -> list[tuple[Mode, Mode, float]]:
    """
    按本征矢重叠配对。总 beta 对模式排列不敏感，但模式贡献和频段贡献需要配对。
    """
    if not modes6 or not modes7:
        return []

    n6 = len(modes6)
    n7 = len(modes7)
    overlap = np.zeros((n6, n7), dtype=float)
    for i, mode6 in enumerate(modes6):
        for j, mode7 in enumerate(modes7):
            overlap[i, j] = vector_overlap(mode6, mode7)

    pairs: list[tuple[int, int]] = []

    if linear_sum_assignment is not None:
        rows, cols = linear_sum_assignment(-overlap)
        pairs = list(zip(rows.tolist(), cols.tolist()))
    else:
        # 没有 scipy 时使用贪心最大重叠匹配。
        candidates = [
            (float(overlap[i, j]), i, j)
            for i in range(n6)
            for j in range(n7)
        ]
        candidates.sort(reverse=True)
        used_i: set[int] = set()
        used_j: set[int] = set()
        for _, i, j in candidates:
            if i in used_i or j in used_j:
                continue
            pairs.append((i, j))
            used_i.add(i)
            used_j.add(j)
            if len(pairs) == min(n6, n7):
                break

    result = [
        (modes6[i], modes7[j], float(overlap[i, j]))
        for i, j in pairs
    ]
    result.sort(key=lambda item: item[0].index)
    return result


# -----------------------------
# beta^(7/6) 计算
# -----------------------------
def log_sinh(x: float) -> float:
    if x <= 0:
        raise ValueError(f"log_sinh 要求 x>0，得到 {x}")
    if x < 20.0:
        return math.log(math.sinh(x))
    return x - math.log(2.0) + math.log1p(-math.exp(-2.0 * x))


def mode_ln_beta_7_6(freq6_thz: float, freq7_thz: float, temperature: float) -> float:
    """
    beta^(7/6) 的单模对数贡献：
      ln beta_i = ln(x7/x6) + ln[sinh(x6)] - ln[sinh(x7)]
      x_s = h nu_s / (2 k_B T)
    """
    if freq6_thz <= 0 or freq7_thz <= 0:
        raise ValueError("beta 计算要求两套频率均为正实频。")

    x6 = H_EV_S * freq6_thz * 1e12 / (2.0 * KB_EV_K * temperature)
    x7 = H_EV_S * freq7_thz * 1e12 / (2.0 * KB_EV_K * temperature)

    return math.log(x7 / x6) + log_sinh(x6) - log_sinh(x7)


def li_participation(mode: Mode, atom_symbols: list[str]) -> float | None:
    if mode.eigenvector is None:
        return None
    if len(atom_symbols) != mode.eigenvector.shape[0]:
        return None

    amplitudes = np.sum(mode.eigenvector**2, axis=1)
    total = float(np.sum(amplitudes))
    if total <= 0:
        return None

    li_sum = float(
        sum(amplitude for amplitude, symbol in zip(amplitudes, atom_symbols) if symbol == "Li")
    )
    return li_sum / total


def calculate_material_beta(
    material_id: str,
    freq6_path: Path,
    freq7_path: Path,
    structure_info: StructureInfo,
    temperature: float,
    acoustic_cutoff_cm1: float,
) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    modes6_all = parse_frequency_file(freq6_path)
    modes7_all = parse_frequency_file(freq7_path)

    # 先对全部模式进行匹配，再整对排除虚频。这样可以避免
    # MASS=6 中为虚频、MASS=7 中为实频的同一模式错配到别的模式。
    all_pairs = pair_modes(modes6_all, modes7_all)
    pairs, excluded_imaginary_pairs, excluded_low_frequency_pairs = (
        select_usable_pairs(all_pairs, acoustic_cutoff_cm1)
    )

    imag6 = imaginary_modes(modes6_all)
    imag7 = imaginary_modes(modes7_all)

    notes: list[str] = []
    if len(modes6_all) != 3 * structure_info.n_atoms:
        notes.append(
            f"MASS=6模式数{len(modes6_all)}不等于3N={3 * structure_info.n_atoms}"
        )
    if len(modes7_all) != 3 * structure_info.n_atoms:
        notes.append(
            f"MASS=7模式数{len(modes7_all)}不等于3N={3 * structure_info.n_atoms}"
        )
    if imag6:
        notes.append(f"MASS=6的{len(imag6)}个虚频已全部排除")
    if imag7:
        notes.append(f"MASS=7的{len(imag7)}个虚频已全部排除")
    if excluded_imaginary_pairs:
        notes.append(f"因至少一侧为虚频而排除{excluded_imaginary_pairs}个模式对")
    if excluded_low_frequency_pairs:
        notes.append(
            f"因两侧任一频率不高于{acoustic_cutoff_cm1:g} cm^-1"
            f"而排除{excluded_low_frequency_pairs}个模式对"
        )
    if len(all_pairs) != min(len(modes6_all), len(modes7_all)):
        notes.append("全部模式配对数量异常")

    mode_rows: list[dict[str, Any]] = []
    ln_beta_values: list[float] = []

    for pair_number, (mode6, mode7, overlap) in enumerate(pairs, start=1):
        ln_beta = mode_ln_beta_7_6(
            mode6.freq_thz, mode7.freq_thz, temperature
        )
        ln_beta_values.append(ln_beta)

        mean_cm1 = 0.5 * (mode6.freq_cm1 + mode7.freq_cm1)
        mode_rows.append(
            {
                "材料ID": material_id,
                "配对序号": pair_number,
                "MASS6模式编号": mode6.index,
                "MASS7模式编号": mode7.index,
                "MASS6频率_THz": mode6.freq_thz,
                "MASS7频率_THz": mode7.freq_thz,
                "MASS6频率_cm-1": mode6.freq_cm1,
                "MASS7频率_cm-1": mode7.freq_cm1,
                "平均频率_cm-1": mean_cm1,
                "本征矢重叠": overlap,
                "Li参与率_MASS6": li_participation(mode6, structure_info.atom_symbols),
                "Li参与率_MASS7": li_participation(mode7, structure_info.atom_symbols),
                "lnBeta_7_6_单模": ln_beta,
                "1000lnBeta_7_6_单模": 1000.0 * ln_beta,
                "是否300到500cm-1": 300.0 <= mean_cm1 <= 500.0,
            }
        )

    ln_beta_cell = float(sum(ln_beta_values)) if ln_beta_values else math.nan
    ln_beta_per_li = (
        ln_beta_cell / structure_info.n_li if math.isfinite(ln_beta_cell) else math.nan
    )

    if math.isfinite(ln_beta_cell):
        beta_cell = math.exp(ln_beta_cell)
        beta_per_li = math.exp(ln_beta_per_li)
    else:
        beta_cell = math.nan
        beta_per_li = math.nan

    if ln_beta_values and abs(ln_beta_cell) > 1e-15:
        interval_sum = sum(
            row["lnBeta_7_6_单模"]
            for row in mode_rows
            if row["是否300到500cm-1"]
        )
        contrib_300_500 = 100.0 * interval_sum / ln_beta_cell
        for row in mode_rows:
            row["总Beta贡献_pct"] = (
                100.0 * row["lnBeta_7_6_单模"] / ln_beta_cell
            )
    else:
        contrib_300_500 = math.nan
        for row in mode_rows:
            row["总Beta贡献_pct"] = math.nan

    overlaps = [row["本征矢重叠"] for row in mode_rows]
    mean_overlap = float(np.mean(overlaps)) if overlaps else math.nan
    min_overlap = float(np.min(overlaps)) if overlaps else math.nan

    if not mode_rows:
        status = "FAILED_NO_USABLE_REAL_MODE_PAIRS"
        valid_for_correlation = False
    else:
        # 用户要求：全部虚频不参与 beta，但不因存在虚频而否定该材料。
        status = "OK"
        valid_for_correlation = True

    summary = {
        "材料ID": material_id,
        "Beta状态": status,
        "温度_K": temperature,
        "结构文件": str(structure_info.path),
        "MASS6频率文件": str(freq6_path),
        "MASS7频率文件": str(freq7_path),
        "原子数": structure_info.n_atoms,
        "Li原子数": structure_info.n_li,
        "MASS6总模式数": len(modes6_all),
        "MASS7总模式数": len(modes7_all),
        "MASS6实频总数": sum(not mode.imaginary for mode in modes6_all),
        "MASS7实频总数": sum(not mode.imaginary for mode in modes7_all),
        "实际用于Beta的配对模式数": len(mode_rows),
        "MASS6总虚频数": len(imag6),
        "MASS7总虚频数": len(imag7),
        "因虚频排除的配对数": excluded_imaginary_pairs,
        "因声学低频排除的配对数": excluded_low_frequency_pairs,
        "MASS6最大虚频_cm-1": max((mode.freq_cm1 for mode in imag6), default=0.0),
        "MASS7最大虚频_cm-1": max((mode.freq_cm1 for mode in imag7), default=0.0),
        "虚频处理方式": "全部排除；只使用两套中均为正实频的配对模式",
        "Beta_7_6_Gamma_cell": beta_cell,
        "1000lnBeta_7_6_Gamma_cell": 1000.0 * ln_beta_cell,
        "Beta_7_6_Gamma_perLi": beta_per_li,
        "1000lnBeta_7_6_Gamma_perLi": 1000.0 * ln_beta_per_li,
        "300到500cm-1贡献_pct": contrib_300_500,
        "平均本征矢重叠": mean_overlap,
        "最小本征矢重叠": min_overlap,
        "可用于bottleneck相关性": valid_for_correlation,
        "备注": "；".join(notes) if notes else None,
    }

    return summary, mode_rows


# -----------------------------
# 统计相关性
# -----------------------------
def rankdata_average(values: np.ndarray) -> np.ndarray:
    order = np.argsort(values, kind="mergesort")
    ranks = np.empty(len(values), dtype=float)
    i = 0
    while i < len(values):
        j = i + 1
        while j < len(values) and values[order[j]] == values[order[i]]:
            j += 1
        average_rank = 0.5 * ((i + 1) + j)
        ranks[order[i:j]] = average_rank
        i = j
    return ranks


def pearson_correlation(x: list[float], y: list[float]) -> float | None:
    if len(x) < 3 or len(y) != len(x):
        return None
    x_array = np.asarray(x, dtype=float)
    y_array = np.asarray(y, dtype=float)
    if np.std(x_array) <= 0 or np.std(y_array) <= 0:
        return None
    return float(np.corrcoef(x_array, y_array)[0, 1])


def spearman_correlation(x: list[float], y: list[float]) -> float | None:
    if len(x) < 3 or len(y) != len(x):
        return None
    return pearson_correlation(
        rankdata_average(np.asarray(x, dtype=float)).tolist(),
        rankdata_average(np.asarray(y, dtype=float)).tolist(),
    )


# -----------------------------
# Excel 写入
# -----------------------------
BETA_COLUMNS = [
    "Beta状态",
    "温度_K",
    "MASS6频率文件",
    "MASS7频率文件",
    "Li原子数_Beta",
    "Beta配对模式数",
    "MASS6总虚频数",
    "MASS7总虚频数",
    "因虚频排除的配对数",
    "因声学低频排除的配对数",
    "MASS6最大虚频_cm-1",
    "MASS7最大虚频_cm-1",
    "虚频处理方式",
    "Beta_7_6_Gamma_cell",
    "1000lnBeta_7_6_Gamma_cell",
    "Beta_7_6_Gamma_perLi",
    "1000lnBeta_7_6_Gamma_perLi",
    "300到500cm-1贡献_pct",
    "平均本征矢重叠",
    "最小本征矢重叠",
    "可用于bottleneck相关性",
    "Beta备注",
]


def find_header_map(ws) -> dict[str, int]:
    return {
        str(cell.value).strip(): cell.column
        for cell in ws[1]
        if cell.value is not None
    }


def style_new_headers(ws, start_col: int, end_col: int) -> None:
    fill = PatternFill("solid", fgColor="7030A0")
    font = Font(color="FFFFFF", bold=True)
    for col in range(start_col, end_col + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def autofit_worksheet(ws, max_width: int = 36) -> None:
    for column_cells in ws.columns:
        letter = get_column_letter(column_cells[0].column)
        maximum = 8
        for cell in column_cells:
            if cell.value is not None:
                maximum = max(maximum, min(len(str(cell.value)) + 2, max_width))
        ws.column_dimensions[letter].width = maximum


def append_beta_to_material_summary(ws, beta_map: dict[str, dict[str, Any]]) -> None:
    headers = find_header_map(ws)
    material_col = headers.get("材料ID")
    if material_col is None:
        raise ValueError("材料汇总中找不到‘材料ID’列。")

    # 如果列已经存在则复用，否则追加。
    current_headers = find_header_map(ws)
    new_start = ws.max_column + 1
    next_col = new_start
    beta_header_cols: dict[str, int] = {}

    for header in BETA_COLUMNS:
        if header in current_headers:
            beta_header_cols[header] = current_headers[header]
        else:
            ws.cell(row=1, column=next_col, value=header)
            beta_header_cols[header] = next_col
            next_col += 1

    if next_col > new_start:
        style_new_headers(ws, new_start, next_col - 1)

    field_mapping = {
        "Beta状态": "Beta状态",
        "温度_K": "温度_K",
        "MASS6频率文件": "MASS6频率文件",
        "MASS7频率文件": "MASS7频率文件",
        "Li原子数_Beta": "Li原子数",
        "Beta配对模式数": "实际用于Beta的配对模式数",
        "MASS6总虚频数": "MASS6总虚频数",
        "MASS7总虚频数": "MASS7总虚频数",
        "因虚频排除的配对数": "因虚频排除的配对数",
        "因声学低频排除的配对数": "因声学低频排除的配对数",
        "MASS6最大虚频_cm-1": "MASS6最大虚频_cm-1",
        "MASS7最大虚频_cm-1": "MASS7最大虚频_cm-1",
        "虚频处理方式": "虚频处理方式",
        "Beta_7_6_Gamma_cell": "Beta_7_6_Gamma_cell",
        "1000lnBeta_7_6_Gamma_cell": "1000lnBeta_7_6_Gamma_cell",
        "Beta_7_6_Gamma_perLi": "Beta_7_6_Gamma_perLi",
        "1000lnBeta_7_6_Gamma_perLi": "1000lnBeta_7_6_Gamma_perLi",
        "300到500cm-1贡献_pct": "300到500cm-1贡献_pct",
        "平均本征矢重叠": "平均本征矢重叠",
        "最小本征矢重叠": "最小本征矢重叠",
        "可用于bottleneck相关性": "可用于bottleneck相关性",
        "Beta备注": "备注",
    }

    for row in range(2, ws.max_row + 1):
        raw_id = ws.cell(row=row, column=material_col).value
        key = normalized_material_id(raw_id)
        record = beta_map.get(key)

        if record is None:
            ws.cell(row=row, column=beta_header_cols["Beta状态"], value="MISSING_FREQUENCY_PAIR")
            ws.cell(
                row=row,
                column=beta_header_cols["Beta备注"],
                value="没有找到 MASS=6 和 MASS=7 的完整频率配对。",
            )
            ws.cell(
                row=row,
                column=beta_header_cols["可用于bottleneck相关性"],
                value=False,
            )
            continue

        for excel_header, record_key in field_mapping.items():
            value = record.get(record_key)
            if isinstance(value, Path):
                value = str(value)
            ws.cell(row=row, column=beta_header_cols[excel_header], value=value)

    # 数字格式
    for header in [
        "Beta_7_6_Gamma_cell",
        "Beta_7_6_Gamma_perLi",
        "平均本征矢重叠",
        "最小本征矢重叠",
    ]:
        col = beta_header_cols[header]
        for row in range(2, ws.max_row + 1):
            ws.cell(row=row, column=col).number_format = "0.000000"

    for header in [
        "1000lnBeta_7_6_Gamma_cell",
        "1000lnBeta_7_6_Gamma_perLi",
        "300到500cm-1贡献_pct",
        "MASS6最大虚频_cm-1",
        "MASS7最大虚频_cm-1",
    ]:
        col = beta_header_cols[header]
        for row in range(2, ws.max_row + 1):
            ws.cell(row=row, column=col).number_format = "0.000"

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def write_table_sheet(ws, records: list[dict[str, Any]], headers: list[str]) -> None:
    ws.append(headers)
    for record in records:
        ws.append([record.get(header) for header in headers])

    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    autofit_worksheet(ws)


def replace_sheet(workbook, title: str):
    if title in workbook.sheetnames:
        workbook.remove(workbook[title])
    return workbook.create_sheet(title)


def build_correlation_sheet(
    workbook,
    material_ws,
    beta_map: dict[str, dict[str, Any]],
) -> None:
    ws = replace_sheet(workbook, "bottleneck_beta关联")
    material_headers = find_header_map(material_ws)

    required = ["材料ID", "化学式", "推荐_bottleneck半径_A", "几何贯通维数_CAVD"]
    missing = [header for header in required if header not in material_headers]
    if missing:
        raise ValueError(f"材料汇总缺少关联分析列：{missing}")

    headers = [
        "材料ID",
        "化学式",
        "几何贯通维数_CAVD",
        "Li几何贯通判定",
        "推荐_bottleneck半径_A",
        "推荐_bottleneck直径_A",
        "Beta状态",
        "Beta_7_6_Gamma_perLi",
        "1000lnBeta_7_6_Gamma_perLi",
        "300到500cm-1贡献_pct",
        "可用于相关性",
        "纳入贯通材料相关性",
        "备注",
    ]

    rows: list[dict[str, Any]] = []
    x_all: list[float] = []
    y_all: list[float] = []
    x_connected: list[float] = []
    y_connected: list[float] = []

    for row_index in range(2, material_ws.max_row + 1):
        material_id = material_ws.cell(
            row=row_index, column=material_headers["材料ID"]
        ).value
        key = normalized_material_id(material_id)
        beta = beta_map.get(key)

        bottleneck = material_ws.cell(
            row=row_index, column=material_headers["推荐_bottleneck半径_A"]
        ).value
        dimension = material_ws.cell(
            row=row_index, column=material_headers["几何贯通维数_CAVD"]
        ).value

        beta_valid = bool(beta and beta.get("可用于bottleneck相关性"))
        numeric_bottleneck = isinstance(bottleneck, (int, float))
        beta_value = beta.get("1000lnBeta_7_6_Gamma_perLi") if beta else None
        numeric_beta = isinstance(beta_value, (int, float)) and math.isfinite(float(beta_value))
        include_all = beta_valid and numeric_bottleneck and numeric_beta
        include_connected = include_all and isinstance(dimension, (int, float)) and float(dimension) > 0

        if include_all:
            x_all.append(float(bottleneck))
            y_all.append(float(beta_value))
        if include_connected:
            x_connected.append(float(bottleneck))
            y_connected.append(float(beta_value))

        rows.append(
            {
                "材料ID": material_id,
                "化学式": material_ws.cell(
                    row=row_index, column=material_headers["化学式"]
                ).value,
                "几何贯通维数_CAVD": dimension,
                "Li几何贯通判定": material_ws.cell(
                    row=row_index,
                    column=material_headers.get("Li几何贯通判定", material_headers["几何贯通维数_CAVD"]),
                ).value,
                "推荐_bottleneck半径_A": bottleneck,
                "推荐_bottleneck直径_A": material_ws.cell(
                    row=row_index,
                    column=material_headers.get("推荐_bottleneck直径_A", material_headers["推荐_bottleneck半径_A"]),
                ).value,
                "Beta状态": beta.get("Beta状态") if beta else "MISSING_FREQUENCY_PAIR",
                "Beta_7_6_Gamma_perLi": beta.get("Beta_7_6_Gamma_perLi") if beta else None,
                "1000lnBeta_7_6_Gamma_perLi": beta_value,
                "300到500cm-1贡献_pct": beta.get("300到500cm-1贡献_pct") if beta else None,
                "可用于相关性": include_all,
                "纳入贯通材料相关性": include_connected,
                "备注": beta.get("备注") if beta else "缺少完整频率配对",
            }
        )

    # 先写统计摘要，再写明细。
    ws["A1"] = "bottleneck–beta 相关性摘要"
    ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="7030A0")
    ws.merge_cells("A1:D1")

    statistics = [
        ["样本范围", "样本数", "Pearson r", "Spearman rho"],
        [
            "所有Beta有效材料（含0D）",
            len(x_all),
            pearson_correlation(x_all, y_all),
            spearman_correlation(x_all, y_all),
        ],
        [
            "仅几何贯通材料（维数>0）",
            len(x_connected),
            pearson_correlation(x_connected, y_connected),
            spearman_correlation(x_connected, y_connected),
        ],
    ]
    for row in statistics:
        ws.append(row)

    for cell in ws[2]:
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.font = Font(bold=True)

    start_row = 6
    for col, header in enumerate(headers, start=1):
        ws.cell(row=start_row, column=col, value=header)
    for record in rows:
        ws.append([record.get(header) for header in headers])

    for cell in ws[start_row]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.freeze_panes = f"A{start_row + 1}"
    ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(headers))}{ws.max_row}"

    # 条件色阶便于观察趋势。
    bottleneck_col = headers.index("推荐_bottleneck半径_A") + 1
    beta_col = headers.index("1000lnBeta_7_6_Gamma_perLi") + 1
    ws.conditional_formatting.add(
        f"{get_column_letter(bottleneck_col)}{start_row+1}:{get_column_letter(bottleneck_col)}{ws.max_row}",
        ColorScaleRule(start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50, mid_color="FFEB84", end_type="max", end_color="63BE7B"),
    )
    ws.conditional_formatting.add(
        f"{get_column_letter(beta_col)}{start_row+1}:{get_column_letter(beta_col)}{ws.max_row}",
        ColorScaleRule(start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50, mid_color="FFEB84", end_type="max", end_color="63BE7B"),
    )

    autofit_worksheet(ws)


def write_output_workbook(
    input_xlsx: Path,
    output_xlsx: Path,
    beta_records: list[dict[str, Any]],
    mode_records: list[dict[str, Any]],
) -> None:
    output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    if input_xlsx.resolve() != output_xlsx.resolve():
        shutil.copy2(input_xlsx, output_xlsx)

    # 两次读取：
    # 1) 普通模式保留原工作簿公式和样式；
    # 2) data_only=True 读取公式的已缓存数值，用于关联表。
    #    否则 openpyxl 读到的是 '=IF(...)' 字符串，不是 bottleneck 数值。
    workbook = load_workbook(output_xlsx, data_only=False)
    values_workbook = load_workbook(input_xlsx, data_only=True, read_only=False)
    if "材料汇总" not in workbook.sheetnames:
        raise ValueError("输入 Excel 中找不到‘材料汇总’工作表。")
    if "材料汇总" not in values_workbook.sheetnames:
        raise ValueError("输入 Excel 的缓存值工作簿中找不到‘材料汇总’工作表。")

    material_ws = workbook["材料汇总"]
    material_values_ws = values_workbook["材料汇总"]
    beta_map = {
        normalized_material_id(record["材料ID"]): record
        for record in beta_records
    }

    append_beta_to_material_summary(material_ws, beta_map)

    beta_ws = replace_sheet(workbook, "Beta汇总")
    beta_headers = [
        "材料ID",
        "Beta状态",
        "温度_K",
        "原子数",
        "Li原子数",
        "MASS6总模式数",
        "MASS7总模式数",
        "MASS6实频总数",
        "MASS7实频总数",
        "实际用于Beta的配对模式数",
        "MASS6总虚频数",
        "MASS7总虚频数",
        "因虚频排除的配对数",
        "因声学低频排除的配对数",
        "MASS6最大虚频_cm-1",
        "MASS7最大虚频_cm-1",
        "虚频处理方式",
        "Beta_7_6_Gamma_cell",
        "1000lnBeta_7_6_Gamma_cell",
        "Beta_7_6_Gamma_perLi",
        "1000lnBeta_7_6_Gamma_perLi",
        "300到500cm-1贡献_pct",
        "平均本征矢重叠",
        "最小本征矢重叠",
        "可用于bottleneck相关性",
        "结构文件",
        "MASS6频率文件",
        "MASS7频率文件",
        "备注",
    ]
    beta_records_sorted = sorted(
        beta_records, key=lambda record: normalized_material_id(record["材料ID"])
    )
    write_table_sheet(beta_ws, beta_records_sorted, beta_headers)

    detail_ws = replace_sheet(workbook, "Beta模式明细")
    detail_headers = [
        "材料ID",
        "配对序号",
        "MASS6模式编号",
        "MASS7模式编号",
        "MASS6频率_THz",
        "MASS7频率_THz",
        "MASS6频率_cm-1",
        "MASS7频率_cm-1",
        "平均频率_cm-1",
        "本征矢重叠",
        "Li参与率_MASS6",
        "Li参与率_MASS7",
        "lnBeta_7_6_单模",
        "1000lnBeta_7_6_单模",
        "总Beta贡献_pct",
        "是否300到500cm-1",
    ]
    write_table_sheet(detail_ws, mode_records, detail_headers)

    build_correlation_sheet(workbook, material_values_ws, beta_map)

    # 强制 Excel/LibreOffice 打开时重新计算原工作簿中的公式。
    try:
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.calculation.calcMode = "auto"
    except Exception:
        pass

    # 将关键工作表放到前面。
    preferred_order = ["材料汇总", "Beta汇总", "bottleneck_beta关联", "Beta模式明细"]
    sheets = {ws.title: ws for ws in workbook.worksheets}
    new_order = [sheets[name] for name in preferred_order if name in sheets]
    new_order.extend(ws for ws in workbook.worksheets if ws.title not in preferred_order)
    workbook._sheets = new_order

    workbook.save(output_xlsx)
    values_workbook.close()


# -----------------------------
# 批量运行
# -----------------------------
def auto_find_bottleneck_xlsx(root: Path) -> Path:
    candidates = []
    for path in root.glob("*.xlsx"):
        name = path.name.lower()
        if "bottleneck" in name and "beta" not in name and not name.startswith("~$"):
            candidates.append(path)
    if not candidates:
        raise FileNotFoundError(
            "没有自动找到 bottleneck Excel，请用 --bottleneck 明确指定。"
        )
    candidates.sort(key=lambda path: path.stat().st_mtime, reverse=True)
    return candidates[0]


def run(args: argparse.Namespace) -> None:
    root = args.root.resolve()
    mass6_dir = (root / args.mass6_dir).resolve()
    mass7_dir = (root / args.mass7_dir).resolve()
    crystal_dir = (root / args.crystal_dir).resolve()

    bottleneck_xlsx = (
        args.bottleneck.resolve()
        if args.bottleneck is not None
        else auto_find_bottleneck_xlsx(root)
    )
    output_xlsx = (
        args.output.resolve()
        if args.output is not None
        else root / f"cavd_bottleneck_beta_忽略全部虚频_{args.temperature:g}K汇总.xlsx"
    )

    print("=" * 80)
    print("Li 同位素 beta^(7/6) + CAVD bottleneck 合并")
    print("=" * 80)
    print("根目录：", root)
    print("MASS=6：", mass6_dir)
    print("MASS=7：", mass7_dir)
    print("结构目录：", crystal_dir)
    print("bottleneck Excel：", bottleneck_xlsx)
    print("输出 Excel：", output_xlsx)
    print("温度：", args.temperature, "K")
    print("声学零模阈值：", args.acoustic_cutoff_cm1, "cm^-1")
    print("虚频处理：全部排除，不设置虚频大小阈值")
    print("模式匹配：", "scipy Hungarian" if linear_sum_assignment is not None else "贪心重叠匹配")

    mass6_index = build_file_index(mass6_dir, {".txt"})
    mass7_index = build_file_index(mass7_dir, {".txt"})
    structure_index = build_file_index(crystal_dir, {".vasp", ".poscar", ".contcar"})

    all_keys = sorted(set(mass6_index) | set(mass7_index))
    print(f"MASS=6 文件数：{len(mass6_index)}")
    print(f"MASS=7 文件数：{len(mass7_index)}")
    print(f"频率材料并集：{len(all_keys)}")

    beta_records: list[dict[str, Any]] = []
    mode_records: list[dict[str, Any]] = []

    for index, key in enumerate(all_keys, start=1):
        display_id = mass7_index.get(key, mass6_index.get(key)).stem
        freq6 = mass6_index.get(key)
        freq7 = mass7_index.get(key)
        structure_path = structure_index.get(key)

        print(f"[{index:02d}/{len(all_keys):02d}] {display_id}")

        if freq6 is None or freq7 is None or structure_path is None:
            missing = []
            if freq6 is None:
                missing.append("MASS=6")
            if freq7 is None:
                missing.append("MASS=7")
            if structure_path is None:
                missing.append("结构文件")
            beta_records.append(
                {
                    "材料ID": display_id,
                    "Beta状态": "MISSING_INPUT",
                    "温度_K": args.temperature,
                    "可用于bottleneck相关性": False,
                    "MASS6频率文件": str(freq6) if freq6 else None,
                    "MASS7频率文件": str(freq7) if freq7 else None,
                    "结构文件": str(structure_path) if structure_path else None,
                    "备注": "缺少：" + "、".join(missing),
                }
            )
            print("  MISSING：", "、".join(missing))
            continue

        try:
            structure_info = read_vasp_structure_info(structure_path)
            summary, details = calculate_material_beta(
                material_id=display_id,
                freq6_path=freq6,
                freq7_path=freq7,
                structure_info=structure_info,
                temperature=args.temperature,
                acoustic_cutoff_cm1=args.acoustic_cutoff_cm1,
            )
            beta_records.append(summary)
            mode_records.extend(details)
            print(
                f"  {summary['Beta状态']} | "
                f"1000lnBeta_perLi={summary['1000lnBeta_7_6_Gamma_perLi']:.6f} | "
                f"modes={summary['实际用于Beta的配对模式数']} | "
                f"imag6/7={summary['MASS6总虚频数']}/{summary['MASS7总虚频数']}"
            )
        except Exception as exc:
            beta_records.append(
                {
                    "材料ID": display_id,
                    "Beta状态": "FAILED",
                    "温度_K": args.temperature,
                    "可用于bottleneck相关性": False,
                    "MASS6频率文件": str(freq6),
                    "MASS7频率文件": str(freq7),
                    "结构文件": str(structure_path),
                    "备注": f"{type(exc).__name__}: {exc}",
                }
            )
            print(f"  FAILED：{type(exc).__name__}: {exc}")

    write_output_workbook(
        input_xlsx=bottleneck_xlsx,
        output_xlsx=output_xlsx,
        beta_records=beta_records,
        mode_records=mode_records,
    )

    success = sum(record.get("Beta状态") == "OK" for record in beta_records)
    warning = 0
    missing_or_failed = len(beta_records) - success

    print("\n" + "=" * 80)
    print("计算结束")
    print("=" * 80)
    print("OK：", success)
    print("WARNING：", warning)
    print("缺失/失败：", missing_or_failed)
    print("输出：", output_xlsx)
    print("\n说明：")
    print("1. Beta 为 Gamma 点谐振近似 beta^(7/6)。")
    print("2. 1000lnBeta_perLi 用于跨不同 Li 数目晶胞的初筛比较。")
    print("3. 所有虚频均被完全排除；只使用两套中均为正实频的配对模式。")
    print("4. 存在虚频的材料仍可输出 beta，并按用户要求纳入相关性分析。")
    print("5. alpha 需要指定第二个参照相，本脚本不把单相 beta 错写成 alpha。")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="计算 Li 同位素 beta^(7/6)，并合并到 CAVD bottleneck Excel。"
    )
    parser.add_argument("--root", type=Path, default=Path("."), help="数据根目录")
    parser.add_argument("--mass6-dir", default="FREQ_MASS=6", help="MASS=6 频率目录名")
    parser.add_argument("--mass7-dir", default="FREQ_MASS=7", help="MASS=7 频率目录名")
    parser.add_argument("--crystal-dir", default="crystal_structure", help="结构目录名")
    parser.add_argument("--bottleneck", type=Path, default=None, help="bottleneck Excel 路径")
    parser.add_argument("--output", type=Path, default=None, help="输出 Excel 路径")
    parser.add_argument("--temperature", type=float, default=DEFAULT_TEMPERATURE, help="温度/K")
    parser.add_argument(
        "--acoustic-cutoff-cm1",
        type=float,
        default=DEFAULT_ACOUSTIC_CUTOFF_CM1,
        help="视为声学零模并排除的频率阈值/cm^-1",
    )
    return parser.parse_args()


if __name__ == "__main__":
    run(parse_args())
