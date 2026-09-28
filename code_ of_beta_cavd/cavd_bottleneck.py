from __future__ import annotations

import argparse
import csv
import inspect
import json
import os
import re
import shutil
import subprocess
import sys
import time
import traceback
from datetime import datetime
from pathlib import Path
from typing import Any

import cavd
import numpy as np
import pandas as pd
from cavd.cavd_consts import LOWER_THRESHOLD, UPPER_THRESHOLD
from pymatgen.analysis.bond_valence import BVAnalyzer
from pymatgen.core import Structure
from pymatgen.io.cif import CifFile, CifWriter


# ============================================================
# 1. 路径与运行设置
# ============================================================

BASE_DIR = Path(__file__).resolve().parent

# 原始优化后 CIF 目录
INPUT_DIR = BASE_DIR / "optimized_cif"

# 含价态 CIF 及 CAVD 自动生成的 .net/.vasp/.vesta/.resex 文件
OXIDIZED_DIR = BASE_DIR / "optimized_cif_oxi_source"

# 每个材料的终端输出和报错
LOG_DIR = BASE_DIR / "cavd_source_logs"

# 每个材料的 JSON 结果，用于断点续算
WORKER_RESULT_DIR = BASE_DIR / "cavd_source_worker_results"

# 最终输出
OUTPUT_EXCEL = BASE_DIR / "cavd_Radical_Voronoi_价态与bottleneck汇总H__.xlsx"
SUMMARY_CSV = BASE_DIR / "cavd_Radical_Voronoi_材料汇总.csv"
SITE_CSV = BASE_DIR / "cavd_Radical_Voronoi_位点价态半径.csv"

MIGRANT = "Li"
SYMPREC = 0.01

LOWER = float(LOWER_THRESHOLD[MIGRANT])
UPPER = float(UPPER_THRESHOLD[MIGRANT])

# 单个材料允许的最长运行时间，单位秒
TIMEOUT_SECONDS = 3600

# True：已成功计算的材料下次运行时跳过
RESUME_SUCCESSFUL = True

# True：每完成一个材料就刷新 CSV/Excel
SAVE_AFTER_EACH_MATERIAL = True


# ============================================================
# 2. 人工价态修正
#
# 自动流程：
#   1. 优先使用本字典；
#   2. 否则使用 BVAnalyzer；
#   3. 再退回化学计量价态猜测。
#
# 键必须是 CIF 文件名去掉 .cif 后的名称。
# ============================================================

MANUAL_ELEMENT_OXIDATION_STATES: dict[str, dict[str, float]] = {
    "mp-8001": {
        "Li": 1,
        "Al": 3,
        "O": -2,
    },

    # 失败后可按下面格式补充：
    #
    # "mp-6521_LLTO": {
    #     "Li": 1,
    #     "La": 3,
    #     "Ti": 4,
    #     "O": -2,
    # },
}


# ============================================================
# 3. 人工位点半径修正
#
# 默认留空，让源码版 CAVD 自己根据价态和配位环境查半径：
#     rad_dict=None
#
# 某个材料自动半径失败时，可在这里提供完整的：
#     CIF标签 -> 半径/Å
#
# 必须包含 CIF 中的所有标签，包括迁移 Li。
# ============================================================

MANUAL_SITE_RADII: dict[str, dict[str, float]] = {
    # 示例：
    # "mp-8001": {
    #     "Li0": 0.90,
    #     "Al1": 0.675,
    #     "O2": 1.26,
    #     "O3": 1.26,
    # },
}


# ============================================================
# 4. 通用工具
# ============================================================

def json_safe(value: Any) -> Any:
    if isinstance(value, dict):
        return {
            str(key): json_safe(item)
            for key, item in value.items()
        }

    if isinstance(value, (list, tuple, set)):
        return [json_safe(item) for item in value]

    if isinstance(value, np.ndarray):
        return value.tolist()

    if isinstance(value, np.generic):
        return value.item()

    if isinstance(value, Path):
        return str(value)

    return value


def dumps_json(value: Any) -> str:
    return json.dumps(
        json_safe(value),
        ensure_ascii=False,
        sort_keys=True,
    )


def safe_float(value: Any) -> float | None:
    if value is None:
        return None

    try:
        result = float(value)
    except (TypeError, ValueError):
        return None

    if not np.isfinite(result):
        return None

    return result


def write_json_atomic(data: dict[str, Any], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)

    temp_path = path.with_suffix(path.suffix + ".tmp")

    with temp_path.open("w", encoding="utf-8") as handle:
        json.dump(
            json_safe(data),
            handle,
            ensure_ascii=False,
            indent=2,
        )

    temp_path.replace(path)


def load_json(path: Path) -> dict[str, Any]:
    with path.open("r", encoding="utf-8") as handle:
        return json.load(handle)


def calculate_total_charge(structure: Structure) -> float:
    total = 0.0

    for site in structure:
        for specie, occupancy in site.species.items():
            if not hasattr(specie, "oxi_state"):
                raise ValueError(
                    f"物种 {specie} 没有氧化态。"
                )

            total += (
                float(specie.oxi_state)
                * float(occupancy)
            )

    return total


def oxidation_summary(structure: Structure) -> str:
    states: dict[str, set[float]] = {}

    for site in structure:
        for specie in site.species:
            states.setdefault(specie.symbol, set())
            states[specie.symbol].add(
                float(specie.oxi_state)
            )

    parts: list[str] = []

    for symbol in sorted(states):
        state_text = "/".join(
            f"{state:+g}"
            for state in sorted(states[symbol])
        )
        parts.append(f"{symbol}:{state_text}")

    return "; ".join(parts)


def validate_oxidized_structure(
    structure: Structure,
) -> None:
    if not structure.is_ordered:
        raise ValueError(
            "结构包含部分占位或无序位点；"
            "当前脚本只处理有序结构。"
        )

    has_migrant = False

    for site_index, site in enumerate(structure):
        for specie in site.species:
            if not hasattr(specie, "oxi_state"):
                raise ValueError(
                    f"位点 {site_index} 的 {specie} 没有价态。"
                )

            state = float(specie.oxi_state)

            if abs(state - round(state)) > 1e-6:
                raise ValueError(
                    f"位点 {site_index} 的 {specie} "
                    f"为非整数价态 {state:+g}。"
                )

            if specie.symbol == MIGRANT:
                has_migrant = True

                if abs(state - 1.0) > 1e-6:
                    raise ValueError(
                        f"{MIGRANT} 的价态不是 +1："
                        f"位点 {site_index}，{state:+g}。"
                    )

    if not has_migrant:
        raise ValueError(
            f"结构中没有迁移离子 {MIGRANT}。"
        )

    total_charge = calculate_total_charge(structure)

    if abs(total_charge) > 1e-5:
        raise ValueError(
            f"添加价态后总电荷不为零："
            f"{total_charge:+.8f}"
        )


# ============================================================
# 5. 添加价态
# ============================================================

def decorate_structure(
    structure: Structure,
    material_id: str,
) -> tuple[Structure, str, list[str]]:
    warnings: list[str] = []

    # ---------- 人工元素价态 ----------
    if material_id in MANUAL_ELEMENT_OXIDATION_STATES:
        decorated = structure.copy()

        decorated.add_oxidation_state_by_element(
            MANUAL_ELEMENT_OXIDATION_STATES[
                material_id
            ]
        )

        validate_oxidized_structure(decorated)

        return (
            decorated,
            "manual_element_mapping",
            warnings,
        )

    # ---------- BVAnalyzer，可处理部分位点混合价 ----------
    try:
        decorated = (
            BVAnalyzer()
            .get_oxi_state_decorated_structure(
                structure
            )
        )

        validate_oxidized_structure(decorated)

        return decorated, "BVAnalyzer", warnings

    except Exception as exc:
        warnings.append(
            "BVAnalyzer失败："
            f"{type(exc).__name__}: {exc}"
        )

    # ---------- 化学计量猜测 ----------
    try:
        guesses = (
            structure.composition
            .oxi_state_guesses()
        )

        if not guesses:
            raise ValueError(
                "没有得到化学计量价态候选。"
            )

        first_guess = guesses[0]

        mapping: dict[str, float] = {}

        for element, state in first_guess.items():
            symbol = (
                element.symbol
                if hasattr(element, "symbol")
                else str(element)
            )
            mapping[symbol] = float(state)

        decorated = structure.copy()
        decorated.add_oxidation_state_by_element(
            mapping
        )

        validate_oxidized_structure(decorated)

        warnings.append(
            "使用化学计量价态猜测；"
            "正式分析前应人工核对。"
        )

        return (
            decorated,
            "composition_guess",
            warnings,
        )

    except Exception as exc:
        warnings.append(
            "化学计量价态猜测失败："
            f"{type(exc).__name__}: {exc}"
        )

    raise ValueError(
        "无法自动分配可信的整数价态。"
        "请在 MANUAL_ELEMENT_OXIDATION_STATES "
        "中补充该材料。"
        + "；".join(warnings)
    )


# ============================================================
# 6. 从含价态 CIF 读取位点标签
# ============================================================

def listify(value: Any) -> list[Any]:
    if value is None:
        return []

    if isinstance(value, list):
        return value

    return [value]


def read_cif_site_rows(
    cif_file: Path,
    structure: Structure,
    radii: dict[str, float],
    material_id: str,
    radius_source: str,
) -> list[dict[str, Any]]:
    cif = CifFile.from_file(str(cif_file))

    if not cif.data:
        raise ValueError(
            f"CIF 中没有数据块：{cif_file}"
        )

    block = next(iter(cif.data.values()))
    data = block.data

    labels = listify(
        data.get("_atom_site_label")
    )
    symbols = listify(
        data.get("_atom_site_type_symbol")
        or data.get("_atom_site_label")
    )
    occupancies = listify(
        data.get("_atom_site_occupancy")
    )

    if not labels:
        raise ValueError(
            f"CIF 缺少 _atom_site_label："
            f"{cif_file}"
        )

    if not occupancies:
        occupancies = [1.0] * len(labels)

    if len(symbols) != len(labels):
        raise ValueError(
            "CIF 的标签数量与元素类型数量不一致。"
        )

    if len(occupancies) != len(labels):
        occupancies = [1.0] * len(labels)

    if len(labels) != len(structure):
        raise ValueError(
            "含价态 CIF 位点数与结构位点数不一致："
            f"CIF={len(labels)}，Structure={len(structure)}"
        )

    rows: list[dict[str, Any]] = []

    # CifWriter 对当前有序结构保持位点顺序。
    for site_index, (
        label,
        symbol_text,
        occupancy,
        site,
    ) in enumerate(
        zip(
            labels,
            symbols,
            occupancies,
            structure,
        )
    ):
        label = str(label)

        match = re.search(
            r"[A-Z][a-z]?",
            str(symbol_text),
        )
        element = (
            match.group(0)
            if match
            else site.specie.symbol
        )

        radius = radii.get(label)

        rows.append(
            {
                "材料ID": material_id,
                "位点序号_0起始": site_index,
                "CIF标签": label,
                "元素": element,
                "氧化态": float(
                    site.specie.oxi_state
                ),
                "占位率": safe_float(
                    occupancy
                ) or 1.0,
                "离子半径_A": safe_float(radius),
                "半径来源": radius_source,
                "分数坐标_x": float(
                    site.frac_coords[0]
                ),
                "分数坐标_y": float(
                    site.frac_coords[1]
                ),
                "分数坐标_z": float(
                    site.frac_coords[2]
                ),
            }
        )

    return rows


# ============================================================
# 7. 单材料工作进程
#
# 每个材料由独立 Python 子进程运行。
# 即使某个 CAVD 计算崩溃或调用 SystemExit，
# 也不会导致整个批处理终止。
# ============================================================

def worker_result_path(
    material_id: str,
) -> Path:
    return (
        WORKER_RESULT_DIR
        / f"{material_id}.json"
    )


def run_worker(input_file: Path) -> int:
    started = time.time()

    material_id = input_file.stem
    result_path = worker_result_path(
        material_id
    )

    OXIDIZED_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )
    WORKER_RESULT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    oxidized_cif = (
        OXIDIZED_DIR
        / f"{material_id}_oxi.cif"
    )

    result: dict[str, Any] = {
        "材料ID": material_id,
        "原始CIF": str(
            input_file.relative_to(BASE_DIR)
        ),
        "含价态CIF": str(
            oxidized_cif.relative_to(BASE_DIR)
        ),
        "化学式": None,
        "约化化学式": None,
        "原子数": None,
        "价态状态": "FAILED",
        "价态方法": None,
        "价态汇总": None,
        "总电荷": None,
        "价态警告": None,
        "CAVD状态": "NOT_RUN",
        "Radical_Voronoi": True,
        "半径来源": None,
        "半径字典_JSON": None,
        "R_a_A": None,
        "R_b_A": None,
        "R_c_A": None,
        "R_1D_A_三方向最大": None,
        "R_2D_A_三方向中值": None,
        "R_3D_A_三方向最小": None,
        "D_1D_A_临界探针直径": None,
        "D_2D_A_临界探针直径": None,
        "D_3D_A_临界探针直径": None,
        "a方向贯通": None,
        "b方向贯通": None,
        "c方向贯通": None,
        "贯通方向数": None,
        "几何贯通维数_CAVD": None,
        "独立通道维数_JSON": None,
        "通道判断阈值_A": None,
        "Li_lower_A": LOWER,
        "Li_upper_A": UPPER,
        "migrant_min_distance_A": None,
        "自动生成文件_JSON": None,
        "位点结果": [],
        "错误": None,
        "运行秒数": None,
        "完成时间": None,
    }

    try:
        structure = Structure.from_file(
            input_file
        )

        result["化学式"] = (
            structure.composition.formula
        )
        result["约化化学式"] = (
            structure.composition.reduced_formula
        )
        result["原子数"] = len(structure)

        (
            decorated,
            oxidation_method,
            oxidation_warnings,
        ) = decorate_structure(
            structure,
            material_id,
        )

        result["价态状态"] = "OK"
        result["价态方法"] = oxidation_method
        result["价态汇总"] = oxidation_summary(
            decorated
        )
        result["总电荷"] = (
            calculate_total_charge(decorated)
        )
        result["价态警告"] = (
            "；".join(oxidation_warnings)
            if oxidation_warnings
            else None
        )

        CifWriter(
            decorated,
            symprec=None,
            significant_figures=10,
            refine_struct=False,
        ).write_file(oxidized_cif)

        manual_radii = (
            MANUAL_SITE_RADII.get(
                material_id
            )
        )

        radius_source = (
            "manual_site_radii"
            if manual_radii is not None
            else (
                "source_CAVD_auto:"
                "oxidation_state+coordination"
            )
        )

        (
            radii,
            min_radius,
            connection_values,
            connected,
            network_dimension,
            channel_dimensions,
            migrant_min_distance,
        ) = cavd.bmd_com(
            filename=str(oxidized_cif),
            migrant=MIGRANT,
            rad_flag=True,
            lower=LOWER,
            upper=UPPER,
            rad_dict=manual_radii,
            symprec=SYMPREC,
        )

        if len(connection_values) < 3:
            raise ValueError(
                "CAVD 返回的 connection_values "
                f"少于三个方向："
                f"{connection_values}"
            )

        r_a = float(connection_values[0])
        r_b = float(connection_values[1])
        r_c = float(connection_values[2])

        ordered_radii = sorted(
            [r_a, r_b, r_c],
            reverse=True,
        )

        r_1d, r_2d, r_3d = ordered_radii

        connected_values = [
            bool(connected[0]),
            bool(connected[1]),
            bool(connected[2]),
        ]

        radii_dict = {
            str(label): float(radius)
            for label, radius in radii.items()
        }

        generated_files = sorted(
            str(path.relative_to(BASE_DIR))
            for path in OXIDIZED_DIR.glob(
                f"{material_id}_oxi*"
            )
            if path.is_file()
        )

        site_rows = read_cif_site_rows(
            cif_file=oxidized_cif,
            structure=decorated,
            radii=radii_dict,
            material_id=material_id,
            radius_source=radius_source,
        )

        result.update(
            {
                "CAVD状态": "OK",
                "半径来源": radius_source,
                "半径字典_JSON": dumps_json(
                    radii_dict
                ),
                "R_a_A": r_a,
                "R_b_A": r_b,
                "R_c_A": r_c,
                "R_1D_A_三方向最大": r_1d,
                "R_2D_A_三方向中值": r_2d,
                "R_3D_A_三方向最小": r_3d,
                "D_1D_A_临界探针直径":
                    2.0 * r_1d,
                "D_2D_A_临界探针直径":
                    2.0 * r_2d,
                "D_3D_A_临界探针直径":
                    2.0 * r_3d,
                "a方向贯通":
                    connected_values[0],
                "b方向贯通":
                    connected_values[1],
                "c方向贯通":
                    connected_values[2],
                "贯通方向数":
                    sum(connected_values),
                "几何贯通维数_CAVD": (
                    int(network_dimension)
                    if network_dimension is not None
                    else None
                ),
                "独立通道维数_JSON":
                    dumps_json(
                        channel_dimensions
                    ),
                "通道判断阈值_A":
                    safe_float(min_radius),
                "migrant_min_distance_A":
                    safe_float(
                        migrant_min_distance
                    ),
                "自动生成文件_JSON":
                    dumps_json(
                        generated_files
                    ),
                "位点结果": site_rows,
                "错误": None,
            }
        )

    except KeyboardInterrupt:
        raise

    except BaseException as exc:
        result["CAVD状态"] = (
            "FAILED"
            if result["价态状态"] == "OK"
            else "NOT_RUN"
        )

        result["错误"] = (
            f"{type(exc).__name__}: {exc}"
        )

        traceback.print_exc()

    result["运行秒数"] = (
        time.time() - started
    )
    result["完成时间"] = (
        datetime.now().isoformat(
            timespec="seconds"
        )
    )

    write_json_atomic(
        result,
        result_path,
    )

    return (
        0
        if result["CAVD状态"] == "OK"
        else 1
    )


# ============================================================
# 8. 汇总输出
# ============================================================

SUMMARY_COLUMNS = [
    "序号",
    "材料ID",
    "原始CIF",
    "含价态CIF",
    "化学式",
    "约化化学式",
    "原子数",
    "价态状态",
    "价态方法",
    "价态汇总",
    "总电荷",
    "价态警告",
    "CAVD状态",
    "Radical_Voronoi",
    "半径来源",
    "半径字典_JSON",
    "R_a_A",
    "R_b_A",
    "R_c_A",
    "R_1D_A_三方向最大",
    "R_2D_A_三方向中值",
    "R_3D_A_三方向最小",
    "D_1D_A_临界探针直径",
    "D_2D_A_临界探针直径",
    "D_3D_A_临界探针直径",
    "a方向贯通",
    "b方向贯通",
    "c方向贯通",
    "贯通方向数",
    "几何贯通维数_CAVD",
    "独立通道维数_JSON",
    "通道判断阈值_A",
    "Li_lower_A",
    "Li_upper_A",
    "migrant_min_distance_A",
    "自动生成文件_JSON",
    "日志文件",
    "错误",
    "运行秒数",
    "完成时间",
]

SITE_COLUMNS = [
    "材料ID",
    "位点序号_0起始",
    "CIF标签",
    "元素",
    "氧化态",
    "占位率",
    "离子半径_A",
    "半径来源",
    "分数坐标_x",
    "分数坐标_y",
    "分数坐标_z",
]


def collect_existing_results(
    cif_files: list[Path],
) -> tuple[
    list[dict[str, Any]],
    list[dict[str, Any]],
]:
    summary_records: list[dict[str, Any]] = []
    site_records: list[dict[str, Any]] = []

    for index, cif_file in enumerate(
        cif_files,
        start=1,
    ):
        material_id = cif_file.stem
        result_path = worker_result_path(
            material_id
        )

        if not result_path.exists():
            continue

        record = load_json(result_path)

        record["序号"] = index
        record["日志文件"] = str(
            (
                LOG_DIR
                / f"{material_id}.log"
            ).relative_to(BASE_DIR)
        )

        site_rows = record.pop(
            "位点结果",
            [],
        )

        summary_records.append(record)
        site_records.extend(site_rows)

    return summary_records, site_records


def write_csv_files(
    summary_records: list[dict[str, Any]],
    site_records: list[dict[str, Any]],
) -> tuple[pd.DataFrame, pd.DataFrame]:
    summary_df = pd.DataFrame(
        summary_records
    )

    for column in SUMMARY_COLUMNS:
        if column not in summary_df.columns:
            summary_df[column] = None

    summary_df = summary_df[
        SUMMARY_COLUMNS
    ]

    site_df = pd.DataFrame(site_records)

    for column in SITE_COLUMNS:
        if column not in site_df.columns:
            site_df[column] = None

    site_df = site_df[SITE_COLUMNS]

    summary_df.to_csv(
        SUMMARY_CSV,
        index=False,
        encoding="utf-8-sig",
    )
    site_df.to_csv(
        SITE_CSV,
        index=False,
        encoding="utf-8-sig",
    )

    return summary_df, site_df


def get_excel_engine() -> str | None:
    try:
        import openpyxl  # noqa: F401
        return "openpyxl"
    except ImportError:
        pass

    try:
        import xlsxwriter  # noqa: F401
        return "xlsxwriter"
    except ImportError:
        return None


def write_excel(
    summary_df: pd.DataFrame,
    site_df: pd.DataFrame,
) -> bool:
    engine = get_excel_engine()

    if engine is None:
        return False

    info_df = pd.DataFrame(
        [
            [
                "生成时间",
                datetime.now().isoformat(
                    timespec="seconds"
                ),
            ],
            ["CAVD路径", cavd.__file__],
            ["输入目录", str(INPUT_DIR)],
            ["含价态结构目录", str(OXIDIZED_DIR)],
            ["日志目录", str(LOG_DIR)],
            ["迁移离子", MIGRANT],
            ["Li lower / Å", LOWER],
            ["Li upper / Å", UPPER],
            ["symprec", SYMPREC],
            [
                "计算类型",
                "带离子半径的 Radical Voronoi",
            ],
            [
                "R_a/R_b/R_c",
                "沿晶格 a/b/c 方向的临界贯通半径",
            ],
            [
                "R_1D",
                "三个方向中的最大值；"
                "至少一个方向贯通的临界尺度",
            ],
            [
                "R_2D",
                "三个方向的中值；"
                "至少两个方向贯通的临界尺度",
            ],
            [
                "R_3D",
                "三个方向中的最小值；"
                "三个方向均贯通的临界尺度",
            ],
            [
                "重要说明",
                "CAVD 是几何与拓扑描述符，"
                "不等同于 DFT-NEB 迁移能垒。",
            ],
        ],
        columns=["项目", "说明或数值"],
    )

    with pd.ExcelWriter(
        OUTPUT_EXCEL,
        engine=engine,
    ) as writer:
        summary_df.to_excel(
            writer,
            sheet_name="材料汇总",
            index=False,
        )
        site_df.to_excel(
            writer,
            sheet_name="位点价态半径",
            index=False,
        )
        info_df.to_excel(
            writer,
            sheet_name="运行说明",
            index=False,
        )

    # 仅在 openpyxl 可用时进一步美化
    if engine == "openpyxl":
        from openpyxl import load_workbook
        from openpyxl.styles import (
            Alignment,
            Font,
            PatternFill,
        )

        workbook = load_workbook(
            OUTPUT_EXCEL
        )

        header_fill = PatternFill(
            "solid",
            fgColor="1F4E78",
        )
        header_font = Font(
            color="FFFFFF",
            bold=True,
        )
        ok_fill = PatternFill(
            "solid",
            fgColor="C6EFCE",
        )
        fail_fill = PatternFill(
            "solid",
            fgColor="FFC7CE",
        )

        for worksheet in workbook.worksheets:
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = (
                worksheet.dimensions
            )
            worksheet.sheet_view.showGridLines = (
                False
            )

            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                    wrap_text=True,
                )

            for column_cells in worksheet.columns:
                letter = (
                    column_cells[0]
                    .column_letter
                )

                max_length = 10

                for cell in column_cells:
                    if cell.value is None:
                        continue

                    max_length = max(
                        max_length,
                        min(
                            len(str(cell.value)),
                            45,
                        ),
                    )

                worksheet.column_dimensions[
                    letter
                ].width = min(
                    max_length + 2,
                    45,
                )

        worksheet = workbook["材料汇总"]

        headers = {
            cell.value: cell.column
            for cell in worksheet[1]
        }

        status_column = headers.get(
            "CAVD状态"
        )

        if status_column:
            for row_index in range(
                2,
                worksheet.max_row + 1,
            ):
                cell = worksheet.cell(
                    row=row_index,
                    column=status_column,
                )

                if cell.value == "OK":
                    cell.fill = ok_fill
                else:
                    cell.fill = fail_fill

        workbook.save(OUTPUT_EXCEL)

    return True


def save_all_outputs(
    cif_files: list[Path],
) -> None:
    (
        summary_records,
        site_records,
    ) = collect_existing_results(
        cif_files
    )

    (
        summary_df,
        site_df,
    ) = write_csv_files(
        summary_records,
        site_records,
    )

    excel_written = write_excel(
        summary_df,
        site_df,
    )

    if not excel_written:
        print(
            "\n[提示] 没有检测到 openpyxl 或 "
            "xlsxwriter，已正常保存两个 CSV。"
        )
        print(
            "需要 Excel 时安装小型依赖：\n"
            "python -m pip install openpyxl"
        )


# ============================================================
# 9. 主批处理
# ============================================================

def run_batch() -> None:
    print("=" * 80)
    print(
        "源码版 CAVD：批量 Radical Voronoi "
        "bottleneck"
    )
    print("=" * 80)

    print("Python：", sys.version.split()[0])
    print("CAVD：", cavd.__file__)
    print("输入目录：", INPUT_DIR)
    print("Li lower：", LOWER, "Å")
    print("Li upper：", UPPER, "Å")
    print()

    cavd_path = str(
        Path(cavd.__file__).resolve()
    )

    if "cavd_source" not in cavd_path:
        raise RuntimeError(
            "当前导入的似乎不是源码版 CAVD："
            f"{cavd_path}\n"
            "请先 conda activate cavd_src_env。"
        )

    if not INPUT_DIR.exists():
        raise FileNotFoundError(
            f"找不到输入目录：{INPUT_DIR}"
        )

    OXIDIZED_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )
    LOG_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )
    WORKER_RESULT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    cif_files = sorted(
        path
        for path in INPUT_DIR.glob("*.cif")
        if not path.name.endswith(
            "_oxi.cif"
        )
    )

    if not cif_files:
        raise FileNotFoundError(
            f"{INPUT_DIR} 中没有 .cif 文件。"
        )

    print(
        f"发现 CIF 数量：{len(cif_files)}"
    )

    script_path = Path(
        __file__
    ).resolve()

    for index, cif_file in enumerate(
        cif_files,
        start=1,
    ):
        material_id = cif_file.stem
        result_path = worker_result_path(
            material_id
        )
        log_file = (
            LOG_DIR
            / f"{material_id}.log"
        )

        if (
            RESUME_SUCCESSFUL
            and result_path.exists()
        ):
            try:
                old_result = load_json(
                    result_path
                )
                if (
                    old_result.get(
                        "CAVD状态"
                    )
                    == "OK"
                ):
                    print(
                        f"[{index:02d}/"
                        f"{len(cif_files):02d}] "
                        f"{material_id}："
                        "已有成功结果，跳过"
                    )
                    continue
            except Exception:
                pass

        print(
            f"[{index:02d}/"
            f"{len(cif_files):02d}] "
            f"{material_id}"
        )

        # 删除旧结果，防止子进程失败后误读旧 JSON
        result_path.unlink(
            missing_ok=True
        )

        started = time.time()

        with log_file.open(
            "w",
            encoding="utf-8",
        ) as log_handle:
            try:
                completed = subprocess.run(
                    [
                        sys.executable,
                        str(script_path),
                        "--worker",
                        str(cif_file),
                    ],
                    cwd=str(BASE_DIR),
                    stdout=log_handle,
                    stderr=subprocess.STDOUT,
                    timeout=TIMEOUT_SECONDS,
                    check=False,
                    env=os.environ.copy(),
                )

                return_code = (
                    completed.returncode
                )

            except subprocess.TimeoutExpired:
                return_code = 124
                log_handle.write(
                    "\n\n"
                    f"TIMEOUT：超过 "
                    f"{TIMEOUT_SECONDS} 秒。\n"
                )

        elapsed = time.time() - started

        if result_path.exists():
            record = load_json(
                result_path
            )

            if (
                record.get("CAVD状态")
                == "OK"
            ):
                print(
                    "  OK | "
                    f"R=({record['R_a_A']:.6f}, "
                    f"{record['R_b_A']:.6f}, "
                    f"{record['R_c_A']:.6f}) Å | "
                    f"维数="
                    f"{record['几何贯通维数_CAVD']} | "
                    f"通道="
                    f"{record['独立通道维数_JSON']} | "
                    f"{elapsed:.1f}s"
                )
            else:
                print(
                    "  FAILED | ",
                    record.get("错误"),
                )
                print(
                    "  日志：",
                    log_file.relative_to(
                        BASE_DIR
                    ),
                )

        else:
            # 子进程完全崩溃或超时，没有结果 JSON
            failure_record = {
                "材料ID": material_id,
                "原始CIF": str(
                    cif_file.relative_to(
                        BASE_DIR
                    )
                ),
                "CAVD状态": "FAILED",
                "Radical_Voronoi": True,
                "Li_lower_A": LOWER,
                "Li_upper_A": UPPER,
                "错误": (
                    "子进程未生成结果 JSON；"
                    f"return_code={return_code}。"
                    "请查看日志。"
                ),
                "运行秒数": elapsed,
                "完成时间": (
                    datetime.now().isoformat(
                        timespec="seconds"
                    )
                ),
                "位点结果": [],
            }

            write_json_atomic(
                failure_record,
                result_path,
            )

            print(
                "  FAILED | 子进程未生成结果 | "
                f"return_code={return_code}"
            )
            print(
                "  日志：",
                log_file.relative_to(
                    BASE_DIR
                ),
            )

        if SAVE_AFTER_EACH_MATERIAL:
            save_all_outputs(
                cif_files
            )

    save_all_outputs(cif_files)

    summary_records, _ = (
        collect_existing_results(
            cif_files
        )
    )

    success_count = sum(
        record.get("CAVD状态")
        == "OK"
        for record in summary_records
    )

    failed_count = (
        len(cif_files) - success_count
    )

    print("\n" + "=" * 80)
    print("全部处理结束")
    print("=" * 80)
    print("材料总数：", len(cif_files))
    print("成功：", success_count)
    print("未成功：", failed_count)
    print("材料汇总 CSV：", SUMMARY_CSV)
    print("位点结果 CSV：", SITE_CSV)

    if get_excel_engine() is not None:
        print("Excel：", OUTPUT_EXCEL)
    else:
        print(
            "Excel 未生成：缺少 openpyxl/"
            "xlsxwriter，CSV 已正常生成。"
        )

    print("含价态 CIF：", OXIDIZED_DIR)
    print("详细日志：", LOG_DIR)


# ============================================================
# 10. 命令行入口
# ============================================================

def parse_arguments() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "源码版 CAVD 批量 Radical Voronoi "
            "bottleneck 计算"
        )
    )

    parser.add_argument(
        "--worker",
        type=Path,
        help=argparse.SUPPRESS,
    )

    return parser.parse_args()


def main() -> None:
    arguments = parse_arguments()

    if arguments.worker is not None:
        input_file = (
            arguments.worker.resolve()
        )
        raise SystemExit(
            run_worker(input_file)
        )

    run_batch()


if __name__ == "__main__":
    main()
